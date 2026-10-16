"""
scripts/build_vku_report.py
Bộ sinh Báo cáo Học thuật VKU (docs/BaoCao.docx và docs/BaoCao.pdf) cho Đề tài 17:
- Tham khảo quy cách Mẫu 3 - Quyết định 675/QĐ-ĐHVH của VKU cho báo cáo học phần:
  + Khổ giấy A4 (21.0 x 29.7 cm), lề trái 3.0 cm, trên/dưới/phải 2.0 cm.
  + Font chữ Times New Roman 13pt, màu đen (#000000), giãn dòng Multiple 1.3, after 6pt, căn đều.
  + Heading 1/2/3 chuẩn của Word, đánh số chương/mục nhất quán, tiêu đề chương bắt đầu trang mới.
  + Trang bìa trang trọng (không ảnh 3D, không số trang), phần đầu đánh số La Mã (i, ii...),
    phần nội dung đánh số Ả Rập bắt đầu từ 1.
  + Mục lục tự động (TOC field), Danh mục hình vẽ, Danh mục bảng biểu, Danh mục từ viết tắt.
  + Bảng biểu nét mảnh, lặp dòng tiêu đề khi qua trang (tblHeader), không tự ngắt dòng bảng (cantSplit).
- Đọc thông tin nhóm từ nguồn duy nhất docs/project_metadata.json và số liệu kiểm thử từ
  evidence/test_results.json.
- Sinh file tạm trước, cập nhật toàn bộ trường TOC/ số trang và xuất PDF bằng Microsoft Word COM;
  chỉ ghi đè docs/BaoCao.docx và docs/BaoCao.pdf sau khi xuất thành công.
"""

from __future__ import annotations

import shutil
import sys
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

from metadata_loader import load_project_metadata, load_test_results, member_role  # noqa: E402
from generate_diagrams import build_diagrams

DOCS_DIR = PROJECT_ROOT / "docs"
DIAG_DIR = DOCS_DIR / "assets" / "diagrams"
EVIDENCE_DIR = PROJECT_ROOT / "evidence"
BUILD_DIR = PROJECT_ROOT / "build"

BLACK = RGBColor(0x00, 0x00, 0x00)
DARK_GRAY = RGBColor(0x22, 0x22, 0x22)


def set_east_asian_font(style_or_run, font_name: str = "Times New Roman") -> None:
    rpr = style_or_run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    rfonts.set(qn("w:cs"), font_name)
    rfonts.set(qn("w:eastAsia"), font_name)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)
    normal.font.color.rgb = BLACK
    set_east_asian_font(normal, "Times New Roman")
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Times New Roman"
    h1.font.size = Pt(15)
    h1.font.bold = True
    h1.font.color.rgb = BLACK
    set_east_asian_font(h1, "Times New Roman")
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(6)
    h1.paragraph_format.space_after = Pt(12)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Times New Roman"
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = BLACK
    set_east_asian_font(h2, "Times New Roman")
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    h3.font.name = "Times New Roman"
    h3.font.size = Pt(13)
    h3.font.bold = True
    h3.font.italic = True
    h3.font.color.rgb = BLACK
    set_east_asian_font(h3, "Times New Roman")
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3.paragraph_format.space_before = Pt(6)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    for name in ("Figure Caption", "Table Caption"):
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.color.rgb = BLACK
        style.font.size = Pt(11.5)
    for name in ("TOC 1", "TOC 2", "TOC 3"):
        if name in doc.styles:
            doc.styles[name].font.color.rgb = BLACK


def apply_a4_margins(section) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)


def set_section_page_number_format(section, fmt: str = "decimal", start: int | None = 1) -> None:
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn("w:pgNumType"))
    if pgNumType is None:
        pgNumType = OxmlElement("w:pgNumType")
        sectPr.append(pgNumType)
    pgNumType.set(qn("w:fmt"), fmt)
    if start is not None:
        pgNumType.set(qn("w:start"), str(start))


def add_field_run(p, field_text: str, font_size: float = 11.0, bold: bool = False) -> None:
    run_begin = p.add_run()
    run_begin.font.name = "Times New Roman"
    run_begin.font.size = Pt(font_size)
    run_begin.bold = bold
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    run_begin._r.append(fldChar1)

    run_instr = p.add_run()
    run_instr.font.name = "Times New Roman"
    run_instr.font.size = Pt(font_size)
    run_instr.bold = bold
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = f" {field_text} "
    run_instr._r.append(instrText)

    run_sep = p.add_run()
    run_sep.font.name = "Times New Roman"
    run_sep.font.size = Pt(font_size)
    run_sep.bold = bold
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "separate")
    run_sep._r.append(fldChar2)

    run_val = p.add_run("1")
    run_val.font.name = "Times New Roman"
    run_val.font.size = Pt(font_size)
    run_val.bold = bold

    run_end = p.add_run()
    run_end.font.name = "Times New Roman"
    run_end.font.size = Pt(font_size)
    run_end.bold = bold
    fldChar3 = OxmlElement("w:fldChar")
    fldChar3.set(qn("w:fldCharType"), "end")
    run_end._r.append(fldChar3)


def configure_header_footer(section, header_text: str | None, show_page_number: bool) -> None:
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = ""
    if header_text:
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.paragraph_format.space_after = Pt(2)
        r = hp.add_run(header_text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10.5)
        r.italic = True
        r.font.color.rgb = BLACK

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.text = ""
    if show_page_number:
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.paragraph_format.space_before = Pt(2)
        add_field_run(fp, "PAGE", font_size=11.0, bold=False)


def set_cell_shading(cell, color_hex: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=110, right=110) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m_name, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m_name}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_row_cant_split(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def set_row_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))


class ReportBuilder:
    def __init__(self, doc: Document):
        self.doc = doc
        self.figures: list[str] = []
        self.tables: list[str] = []

    def add_chapter_heading(self, text: str, new_page: bool = True) -> None:
        if new_page:
            self.doc.add_page_break()
        p = self.doc.add_paragraph(style="Heading 1")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(15)
        r.font.color.rgb = BLACK

    def add_h2(self, text: str) -> None:
        p = self.doc.add_paragraph(style="Heading 2")
        r = p.add_run(text)
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(14)
        r.font.color.rgb = BLACK

    def add_h3(self, text: str) -> None:
        p = self.doc.add_paragraph(style="Heading 3")
        r = p.add_run(text)
        r.bold = True
        r.italic = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        r.font.color.rgb = BLACK

    def add_para(self, text: str, bold_prefix: str = "") -> None:
        p = self.doc.add_paragraph(style="Normal")
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
            rb.font.name = "Times New Roman"
            rb.font.size = Pt(13)
            rb.font.color.rgb = BLACK
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        r.font.color.rgb = BLACK

    def add_bullet(self, bold_prefix: str, text: str) -> None:
        p = self.doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.paragraph_format.space_after = Pt(4)
        rb = p.add_run(f"• {bold_prefix}")
        rb.bold = True
        rb.font.name = "Times New Roman"
        rb.font.size = Pt(13)
        rb.font.color.rgb = BLACK
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        r.font.color.rgb = BLACK

    def add_code_snippet(self, title: str, code_text: str) -> None:
        p_label = self.doc.add_paragraph(style="Normal")
        p_label.paragraph_format.space_before = Pt(4)
        p_label.paragraph_format.space_after = Pt(2)
        p_label.paragraph_format.keep_with_next = True
        rl = p_label.add_run(title)
        rl.bold = True
        rl.italic = True
        rl.font.name = "Times New Roman"
        rl.font.size = Pt(11.5)
        rl.font.color.rgb = BLACK

        p = self.doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.line_spacing = 1.12
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_together = True

        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F8FAFC")
        pPr.append(shd)

        r = p.add_run(code_text.strip("\n"))
        r.font.name = "Consolas"
        r.font.size = Pt(10)
        r.font.color.rgb = DARK_GRAY

    def add_table_with_caption(
        self,
        caption: str,
        headers: list[str],
        rows: list[list[str]],
        col_widths_cm: list[float] | None = None,
    ) -> None:
        self.tables.append(caption)

        # Chú thích bảng nằm PHÍA TRÊN bảng theo đúng quy định
        p_cap = self.doc.add_paragraph(style="Table Caption")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(3)
        p_cap.paragraph_format.keep_with_next = True
        rc = p_cap.add_run(caption)
        rc.bold = True
        rc.font.name = "Times New Roman"
        rc.font.size = Pt(12)
        rc.font.color.rgb = BLACK

        tbl = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Dòng tiêu đề lặp lại khi qua trang và không bị cắt
        set_row_header(tbl.rows[0])
        set_row_cant_split(tbl.rows[0])
        hdr_cells = tbl.rows[0].cells
        for i, h_text in enumerate(headers):
            set_cell_shading(hdr_cells[i], "E2E8F0")
            set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)
            hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(11.5)
            r.font.color.rgb = BLACK

        for r_idx, row_data in enumerate(rows):
            row = tbl.rows[r_idx + 1]
            set_row_cant_split(row)
            row_cells = row.cells
            for c_idx, val in enumerate(row_data):
                set_cell_Margins = set_cell_margins
                set_cell_Margins(row_cells[c_idx], top=70, bottom=70, left=100, right=100)
                row_cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.TOP
                p = row_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(val)
                r.font.name = "Times New Roman"
                r.font.size = Pt(11.5)
                r.font.color.rgb = BLACK

        if col_widths_cm:
            for row in tbl.rows:
                for i, w_cm in enumerate(col_widths_cm):
                    row.cells[i].width = Cm(w_cm)

        p_sp = self.doc.add_paragraph(style="Normal")
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(4)

    def add_figure_with_caption(
        self,
        img_path: Path,
        caption: str,
        width_cm: float = 15.6,
    ) -> None:
        if not img_path.exists():
            raise FileNotFoundError(f"Missing figure file: {img_path}")
        self.figures.append(caption)

        p_img = self.doc.add_paragraph(style="Normal")
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.paragraph_format.keep_with_next = True
        run = p_img.add_run()
        run.add_picture(str(img_path), width=Cm(width_cm))

        # Chú thích hình nằm PHÍA DƯỚI hình theo đúng quy định
        p_cap = self.doc.add_paragraph(style="Figure Caption")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(8)
        rc = p_cap.add_run(caption)
        rc.italic = True
        rc.bold = True
        rc.font.name = "Times New Roman"
        rc.font.size = Pt(11.5)
        rc.font.color.rgb = BLACK


def build_cover_page(doc: Document, meta: dict) -> None:
    sec1 = doc.sections[0]
    apply_a4_margins(sec1)
    sec1.different_first_page_header_footer = True
    configure_header_footer(sec1, header_text=None, show_page_number=False)

    p_uni = doc.add_paragraph(style="Normal")
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_uni.paragraph_format.line_spacing = 1.25
    p_uni.paragraph_format.space_after = Pt(24)
    university_name = meta['university_name'].replace(' VIỆT - HÀN', '\nVIỆT - HÀN')
    r_uni = p_uni.add_run(
        f"{meta['university_parent']}\n"
        f"{university_name}\n"
        f"{meta['faculty_display']}\n"
        "─────────────────────────"
    )
    r_uni.bold = True
    r_uni.font.name = "Times New Roman"
    r_uni.font.size = Pt(13)
    r_uni.font.color.rgb = BLACK

    p_type = doc.add_paragraph(style="Normal")
    p_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_type.paragraph_format.line_spacing = 1.3
    p_type.paragraph_format.space_before = Pt(30)
    p_type.paragraph_format.space_after = Pt(24)
    rt = p_type.add_run(
        "BÁO CÁO ĐỒ ÁN HỌC PHẦN\n"
        f"{meta['course_name'].upper()}"
    )
    rt.bold = True
    rt.font.name = "Times New Roman"
    rt.font.size = Pt(16)
    rt.font.color.rgb = BLACK

    p_topic = doc.add_paragraph(style="Normal")
    p_topic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_topic.paragraph_format.line_spacing = 1.35
    p_topic.paragraph_format.space_before = Pt(12)
    p_topic.paragraph_format.space_after = Pt(36)
    r_top = p_topic.add_run(
        f"{meta['topic_code'].upper()}:\n"
        f"{meta['topic_title_vi'].upper()}\n"
        f"({meta['topic_title']})"
    )
    r_top.bold = True
    r_top.font.name = "Times New Roman"
    r_top.font.size = Pt(15)
    r_top.font.color.rgb = BLACK

    # Khối thông tin Nhóm và Giảng viên hướng dẫn chuẩn học thuật
    info_tbl = doc.add_table(rows=4, cols=2)
    info_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_rows = [
        ("Giảng viên hướng dẫn:", meta["instructor_display"]),
        ("Nhóm thực hiện:", meta["group_display"]),
        ("Học kỳ / Năm học:", meta["academic_term_display"]),
        (
            "Thành viên thực hiện:",
            "\n".join(
                f"{idx}. {m.get('name')} – MSSV: {m.get('student_id')} – Lớp: {m.get('class_name')}"
                for idx, m in enumerate(meta["members"], start=1)
            )
            if meta["members"]
            else "[Chưa cung cấp danh sách thành viên nhóm (Họ tên, MSSV, Lớp)]",
        ),
    ]
    for idx, (label, val) in enumerate(info_rows):
        row = info_tbl.rows[idx]
        set_row_cant_split(row)
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Cm(5.0)
        c1.width = Cm(10.0)
        set_cell_margins(c0, top=60, bottom=60, left=80, right=80)
        set_cell_margins(c1, top=60, bottom=60, left=80, right=80)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(label)
        r0.bold = True
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(13)
        r0.font.color.rgb = BLACK

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(val)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(13)
        r1.font.color.rgb = BLACK

    if meta["is_draft"]:
        p_note = doc.add_paragraph(style="Normal")
        p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_note.paragraph_format.space_before = Pt(18)
        rn = p_note.add_run(
            "Lưu ý bản nháp (Draft): Thông tin danh sách nhóm, MSSV, Lớp, Khoa và GVHD được trích xuất "
            "từ tệp cấu hình docs/project_metadata.json và cần cập nhật trước khi xuất bản nộp cuối."
        )
        rn.italic = True
        rn.font.name = "Times New Roman"
        rn.font.size = Pt(11)
        rn.font.color.rgb = BLACK

    p_loc = doc.add_paragraph(style="Normal")
    p_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_loc.paragraph_format.space_before = Pt(72)
    p_loc.paragraph_format.space_after = Pt(0)
    rloc = p_loc.add_run(f"{meta['location']} – Năm {meta['year']}")
    rloc.bold = True
    rloc.font.name = "Times New Roman"
    rloc.font.size = Pt(13)
    rloc.font.color.rgb = BLACK


def add_toc_placeholder(doc: Document) -> None:
    p_title = doc.add_paragraph(style="Normal")
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(12)
    rt = p_title.add_run("MỤC LỤC")
    rt.bold = True
    rt.font.name = "Times New Roman"
    rt.font.size = Pt(15)
    rt.font.color.rgb = BLACK

    p_toc = doc.add_paragraph(style="Normal")
    add_field_run(p_toc, 'TOC \\o "1-3" \\h \\z \\u', font_size=12.0, bold=False)


def build_report(out_docx: Path, out_pdf: Path) -> None:
    build_diagrams()
    meta = load_project_metadata()
    results = load_test_results()
    reg_suite = results["regular_suite"]
    san_suite = results["sanitizer_suite"]
    env_info = results["environment"]

    doc = Document()
    doc.core_properties.title = "Báo cáo Đề tài 17 — Cài đặt ls với -a, -l, -R, -S"
    doc.core_properties.author = ", ".join(m["name"] for m in meta["members"]) or "Nhóm đồ án — chờ danh sách thành viên"
    doc.core_properties.language = "vi-VN"
    configure_styles(doc)
    rb = ReportBuilder(doc)

    # ==========================================================================
    # SECTION 1: TRANG BÌA (Không đánh số trang)
    # ==========================================================================
    build_cover_page(doc, meta)

    # ==========================================================================
    # SECTION 2: PHẦN ĐẦU (Số trang La Mã i, ii, iii...)
    # ==========================================================================
    sec_front = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_a4_margins(sec_front)
    sec_front.different_first_page_header_footer = False
    set_section_page_number_format(sec_front, fmt="lowerRoman", start=1)
    configure_header_footer(sec_front, header_text=None, show_page_number=True)

    # 1. Mục lục tự động
    add_toc_placeholder(doc)

    # 2. Danh mục từ viết tắt và thuật ngữ
    p_abbr = doc.add_paragraph(style="Normal")
    p_abbr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_abbr.paragraph_format.page_break_before = True
    p_abbr.paragraph_format.space_before = Pt(6)
    p_abbr.paragraph_format.space_after = Pt(12)
    ra = p_abbr.add_run("DANH MỤC TỪ VIẾT TẮT VÀ THUẬT NGỮ")
    ra.bold = True
    ra.font.name = "Times New Roman"
    ra.font.size = Pt(15)
    ra.font.color.rgb = BLACK

    abbr_rows = [
        ["ASan / LSan / UBSan", "AddressSanitizer / LeakSanitizer / UndefinedBehaviorSanitizer", "Bộ công cụ phát hiện lỗi bộ nhớ và hành vi bất định của trình biên dịch GCC/Clang."],
        ["CLI", "Command-Line Interface", "Giao diện dòng lệnh nhận cờ tùy chọn và toán hạng đường dẫn qua argc, argv."],
        ["Directory entry / dentry", "Mục thư mục / bộ đệm tên VFS", "readdir trả struct dirent ở user space; dentry là cấu trúc bộ đệm đường dẫn trong kernel, không đồng nhất hai cấu trúc."],
        ["DrvFs", "Drive File System", "Trình điều khiển gắn kết phân vùng Windows NTFS vào WSL2 (/mnt/c, /mnt/p)."],
        ["EOF", "End Of File", "Trạng thái kết thúc luồng đọc dữ liệu hoặc kết thúc danh mục thư mục."],
        ["FD", "File Descriptor", "Số nguyên không âm đại diện cho tài nguyên tệp hoặc thư mục đang mở của tiến trình."],
        ["GID / UID", "Group ID / User ID", "Định danh số của nhóm sở hữu và người dùng sở hữu đối tượng tệp."],
        ["Inode", "Index Node", "Cấu trúc lưu trữ siêu dữ liệu của tệp trên hệ thống tệp Unix/Linux."],
        ["POSIX", "Portable Operating System Interface", "Bộ tiêu chuẩn IEEE 1003.1 quy định giao diện lập trình hệ điều hành."],
        ["SUID / SGID", "Set User ID / Set Group ID", "Bit đặc biệt của tệp thực thi; SGID trên thư mục còn liên quan kế thừa nhóm. Cách áp dụng phụ thuộc hệ thống tệp và chính sách mount."],
        ["VFS", "Virtual File System", "Lớp trừu tượng hệ thống tệp ảo bên trong nhân Linux."],
        ["WSL2", "Windows Subsystem for Linux 2", "Môi trường chạy nhân Linux thực trên nền tảng ảo hóa nhẹ của Windows."],
    ]
    rb.add_table_with_caption(
        "Bảng 0.1: Giải thích các từ viết tắt và thuật ngữ sử dụng trong báo cáo",
        ["Từ viết tắt", "Thuật ngữ đầy đủ", "Ý nghĩa kỹ thuật trong đồ án"],
        abbr_rows,
        [3.2, 5.3, 7.5],
    )

    # Giữ chỗ cho Danh mục Hình vẽ và Danh mục Bảng biểu (sẽ điền sau khi tạo xong các chương)
    p_lof_title = doc.add_paragraph(style="Normal")
    p_lof_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lof_title.paragraph_format.page_break_before = True
    p_lof_title.paragraph_format.space_before = Pt(6)
    p_lof_title.paragraph_format.space_after = Pt(10)
    r_lof = p_lof_title.add_run("DANH MỤC HÌNH VẼ VÀ BẢNG BIỂU")
    r_lof.bold = True
    r_lof.font.name = "Times New Roman"
    r_lof.font.size = Pt(15)
    r_lof.font.color.rgb = BLACK

    for label, caption_style in (("Danh mục hình vẽ", "Figure Caption"), ("Danh mục bảng biểu", "Table Caption")):
        p_label = doc.add_paragraph(label)
        p_label.runs[0].bold = True
        add_field_run(doc.add_paragraph(), f'TOC \\h \\z \\t "{caption_style},1"', font_size=11.5)

    # ==========================================================================
    # SECTION 3: NỘI DUNG CHÍNH (Số trang Ả Rập bắt đầu từ 1)
    # ==========================================================================
    sec_main = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_a4_margins(sec_main)
    sec_main.different_first_page_header_footer = False
    set_section_page_number_format(sec_main, fmt="decimal", start=1)
    configure_header_footer(
        sec_main,
        header_text="Đồ án Lập trình Hệ thống – Đề tài 17: Cài đặt lệnh ls (-a, -l, -R, -S)",
        show_page_number=True,
    )

    # --------------------------------------------------------------------------
    # MỞ ĐẦU
    # --------------------------------------------------------------------------
    rb.add_chapter_heading("MỞ ĐẦU", new_page=False)

    rb.add_para(
        "Trong các hệ điều hành họ Unix và Linux, tiện ích dòng lệnh ls là công cụ nền tảng để quan sát "
        "không gian tên thư mục và kiểm tra siêu dữ liệu của các đối tượng tệp [1], [2]. Tuy nhiên, đối với "
        "học phần Lập trình Hệ thống, việc gọi hàm thư viện cấp cao như system(\"ls\") chỉ khởi tạo một tiến "
        "trình vỏ (shell) chạy tệp nhị phân có sẵn của hệ điều hành mà không giúp người lập trình tiếp cận "
        "cơ chế quản lý danh mục thư mục, cấu trúc siêu dữ liệu Inode, quản lý bộ nhớ động hay kiểm soát "
        "File Descriptor.",
        "1. Đặt vấn đề và lý do chọn đề tài: ",
    )
    rb.add_para(
        "Đề tài 17 đặt mục tiêu nghiên cứu cơ chế làm việc của hệ thống tệp Linux thông qua giao diện chuẩn "
        "POSIX.1-2008 [1] và tự cài đặt chương trình dòng lệnh myls bằng ngôn ngữ C (chuẩn C11) hỗ trợ bốn "
        "tùy chọn cốt lõi: -a (hiển thị mục ẩn), -l (định dạng chi tiết), -R (duyệt đệ quy cây thư mục con) "
        "và -S (sắp xếp theo kích thước tệp giảm dần), cả khi dùng độc lập lẫn khi kết hợp tự do.",
        "2. Mục tiêu của đồ án: ",
    )
    rb.add_para(
        "Chương trình được giới hạn rõ trong phạm vi một công cụ dòng lệnh chuẩn POSIX: sử dụng trực tiếp "
        "các hàm hệ thống và thư viện C chuẩn (opendir, readdir, closedir, lstat, stat, readlink, getpwuid, "
        "getgrgid, localtime_r, getopt, qsort) [3], [4]; không mở rộng sang giao diện đồ họa, không phụ thuộc "
        "vào các tiện ích ngoài và xác định rõ ranh giới những tính năng đặc thù của GNU ls không nằm trong "
        "phạm vi đề tài.",
        "3. Phạm vi và phương pháp đánh giá: ",
    )

    # Bảng thông tin nhóm & phân công (lấy từ metadata)
    if meta["members"]:
        member_rows = [
            [
                str(idx),
                m.get("name", ""),
                m.get("student_id", ""),
                m.get("class_name", ""),
                member_role(m, meta["task_allocation"]) or "[Chưa cung cấp phân công]",
            ]
            for idx, m in enumerate(meta["members"], start=1)
        ]
        rb.add_table_with_caption(
            "Bảng 0.2: Thông tin thành viên nhóm và phân công thực hiện",
            ["STT", "Họ và tên", "MSSV", "Lớp", "Phân công nhiệm vụ"],
            member_rows, [1.0, 4.1, 2.4, 1.9, 6.6],
        )
    else:
        rb.add_para("Thông tin nhóm chưa được cung cấp. Cần bổ sung đầy đủ họ tên, MSSV, lớp và phân công của tất cả thành viên trước khi nộp; bản thảo này không giả định số lượng hoặc danh tính thành viên.", "4. Thành viên và phân công: ")

    # --------------------------------------------------------------------------
    # CHƯƠNG 1: PHÂN TÍCH YÊU CẦU VÀ HỢP ĐỒNG HÀNH VI
    # --------------------------------------------------------------------------
    rb.add_chapter_heading("CHƯƠNG 1. PHÂN TÍCH YÊU CẦU VÀ HỢP ĐỒNG HÀNH VI", new_page=True)

    rb.add_h2("1.1. Yêu cầu chức năng của bốn tùy chọn (-a, -l, -R, -S)")
    rb.add_para(
        "Chương trình myls nhận các tùy chọn dòng lệnh và không hoặc nhiều toán hạng đường dẫn (path operands) "
        "theo cú pháp chuẩn POSIX [1]: ./myls [-alRS] [file...]. Khi người dùng không truyền đường dẫn nào, "
        "chương trình mặc định liệt kê thư mục làm việc hiện tại (\".\"). Bảng 1.1 tổng hợp yêu cầu chức năng "
        "và ràng buộc kỹ thuật của từng tùy chọn."
    )

    req_rows = [
        [
            "Mặc định (không cờ)",
            "Liệt kê các mục không bắt đầu bằng dấu '.'; in một tên mục trên mỗi dòng.",
            "Sắp xếp tăng dần theo thứ tự byte của tên mục bằng hàm strcmp() (tương đương môi trường LC_ALL=C) [1], [2].",
        ],
        [
            "-a (All entries)",
            "Hiển thị tất cả các mục trong thư mục, bao gồm tệp ẩn, thư mục ẩn và hai mục đặc biệt '.' và '..'.",
            "Tắt bộ lọc kiểm tra ký tự đầu d_name[0] == '.' trong vòng lặp đọc thư mục.",
        ],
        [
            "-l (Long listing)",
            "In dòng 'total <1KiB-blocks>' ở đầu mỗi thư mục; mỗi mục in: chuỗi 10 ký tự loại tệp và quyền, số hard link, tên owner, tên group, kích thước (hoặc major, minor), thời gian sửa đổi và đích symlink.",
            "Cần thu thập toàn bộ mục trước khi in để tính tổng khối và tìm độ rộng lớn nhất của từng cột nhằm căn lề thẳng hàng.",
        ],
        [
            "-R (Recursive)",
            "Sau khi hiển thị toàn bộ mục của thư mục hiện tại, tiếp tục duyệt đệ quy xuống các thư mục con thực sự theo đúng thứ tự đã sắp xếp.",
            "Bắt buộc loại trừ '.' và '..' và không đi vào các symbolic link gặp trong cây thư mục để tránh vòng lặp vô tận [1].",
        ],
        [
            "-S (Sort by size)",
            "Sắp xếp các mục theo kích thước st_size giảm dần (tệp lớn nhất đứng trước).",
            "So sánh kiểu off_t (64 bit trên môi trường x86_64 đã thử) bằng toán tử quan hệ để tránh tràn số; nếu hai mục có cùng st_size thì phân định thứ tự theo tên tăng dần.",
        ],
    ]
    rb.add_table_with_caption(
        "Bảng 1.1: Đặc tả chức năng và ràng buộc kỹ thuật của bốn tùy chọn trong Đề tài 17",
        ["Tùy chọn", "Hành vi hiển thị đầu ra", "Ràng buộc kỹ thuật bắt buộc"],
        req_rows,
        [2.8, 6.6, 6.6],
    )

    rb.add_h2("1.2. Hợp đồng hành vi cho tổ hợp cờ, đối số hỗn hợp và mã thoát")
    rb.add_para(
        "Để chương trình hoạt động nhất quán khi người dùng kết hợp nhiều tùy chọn (như -la, -lS, -aR, -alRS "
        "hoặc viết rời -S -R -l -a) và truyền nhiều đường dẫn hỗn hợp, hợp đồng hành vi của myls được xác định "
        "chặt chẽ qua các quy tắc trong Bảng 1.2."
    )

    contract_rows = [
        [
            "Đa toán hạng hỗn hợp (tệp + thư mục)",
            "Tách thành hai danh sách: nhóm tệp đơn lẻ được sắp xếp và in trước (không có dòng total); nhóm thư mục được sắp xếp và duyệt sau, ngăn cách bằng một dòng trống.",
        ],
        [
            "Tiêu đề tên thư mục (<path>:)",
            "Chỉ in dòng tiêu đề '<path>:' khi bật cờ -R hoặc khi tổng số toán hạng truyền trên dòng lệnh lớn hơn 1.",
        ],
        [
            "Toán hạng CLI là Symbolic link",
            "Nếu không bật -l (hoặc đường dẫn kết thúc bằng dấu '/'), kiểm tra đích bằng stat(): nếu đích là thư mục thì liệt kê nội dung thư mục đích. Nếu bật -l và không có dấu '/' cuối, hiển thị bản thân symlink [1], [2].",
        ],
        [
            "Dấu phân cách '--' và '--help'",
            "Dấu '--' kết thúc phần phân tích cờ để hỗ trợ tên tệp bắt đầu bằng dấu '-'; cờ '--help' in hướng dẫn sử dụng ra stdout và thoát với mã 0.",
        ],
        [
            "Mã thoát 0 (EXIT_LS_OK)",
            "Toàn bộ đối số và thư mục được xử lý thành công, không phát sinh lỗi hệ thống hay lỗi ghi đầu ra.",
        ],
        [
            "Mã thoát 1 (EXIT_LS_MINOR_ERR)",
            "Xảy ra lỗi cục bộ trên một mục hoặc một thư mục con trong quá trình duyệt cây (ví dụ thư mục con trong -R bị khóa quyền đọc chmod 000, hoặc lỗi readdir/readlink/closedir). Chương trình báo lỗi ra stderr và tiếp tục phần còn lại.",
        ],
        [
            "Mã thoát 2 (EXIT_LS_SERIOUS_ERR)",
            "Toán hạng truyền trực tiếp trên CLI không tồn tại hoặc không mở được, tùy chọn dòng lệnh sai, hết bộ nhớ cấp phát động, hoặc lỗi ghi luồng stdout (ví dụ đầy thiết bị /dev/full) [2].",
        ],
        [
            "Tổng hợp nhiều lỗi",
            "Sử dụng quy tắc giữ mức lỗi cao nhất (elevate_status): lỗi mức 1 xảy ra sau không bao giờ làm hạ mức lỗi 2 đã ghi nhận trước đó.",
        ],
    ]
    rb.add_table_with_caption(
        "Bảng 1.2: Hợp đồng hành vi chi tiết của chương trình myls",
        ["Khía cạnh hành vi", "Quy tắc xử lý đã chốt trong chương trình"],
        contract_rows,
        [4.8, 11.2],
    )

    rb.add_h2("1.3. Ranh giới phạm vi và các tính năng ngoài phạm vi đề tài")
    rb.add_para(
        "Chương trình tập trung giải quyết trọn vẹn bốn tùy chọn của Đề tài 17 theo chuẩn POSIX. Các tính năng "
        "mở rộng đặc thù của GNU coreutils không thuộc phạm vi cài đặt và được ghi nhận tường minh gồm: "
        "(1) chế độ dàn nhiều cột theo độ rộng terminal (ioctl TIOCGWINSZ) và tô màu ANSI; "
        "(2) sắp xếp theo quy tắc ngôn ngữ bản địa (locale collation qua strcoll); "
        "(3) hiển thị danh sách kiểm soát truy cập mở rộng ACL (+) hoặc ngữ cảnh bảo mật SELinux (.); "
        "(4) tự động phát hiện chu trình gắn kết hệ thống tệp (mount point cycles) hoặc đồng bộ hóa khi cây "
        "thư mục bị biến đổi đồng thời bởi tiến trình khác trong lúc đang duyệt."
    )

    # --------------------------------------------------------------------------
    # CHƯƠNG 2: CƠ SỞ KỸ THUẬT HỆ THỐNG TỆP LINUX VÀ GIAO DIỆN POSIX
    # --------------------------------------------------------------------------
    rb.add_chapter_heading("CHƯƠNG 2. CƠ SỞ KỸ THUẬT HỆ THỐNG TỆP LINUX VÀ GIAO DIỆN POSIX", new_page=True)

    rb.add_h2("2.1. Cấu trúc danh mục thư mục (Directory Entry) và siêu dữ liệu (Inode)")
    rb.add_para(
        "Trong kiến trúc Hệ thống tệp ảo (Virtual File System – VFS) của nhân Linux [4], tên tệp và siêu dữ liệu "
        "của tệp được lưu trữ ở hai cấu trúc tách biệt:"
    )
    rb.add_bullet(
        "Mục danh bạ thư mục (Directory Entry): ",
        "Trên đĩa, một thư mục là một bảng ánh xạ từ tên mục (d_name) sang số hiệu định danh Inode (d_ino). "
        "Trong nhân Linux, các ánh xạ đường dẫn được lưu đệm tại dentry cache. Khi chương trình ở không gian "
        "người dùng gọi hàm thư viện readdir(3) [3], dữ liệu trả về là cấu trúc struct dirent chứa trường tên "
        "d_name và số d_ino. Cấu trúc struct dirent không chứa kích thước tệp, quyền truy cập, chủ sở hữu hay "
        "dấu thời gian sửa đổi.",
    )
    rb.add_bullet(
        "Nút chỉ mục siêu dữ liệu (Inode và struct stat): ",
        "Mọi thuộc tính vật lý của tệp (loại tệp, 12 bit quyền truy cập, số lượng hard link st_nlink, định danh "
        "chủ sở hữu st_uid, định danh nhóm st_gid, kích thước byte st_size, số hiệu thiết bị st_rdev, số khối "
        "512-byte st_blocks và thời gian sửa đổi st_mtime) đều nằm trong Inode của hệ thống tệp. Để đọc các "
        "thuộc tính này vào không gian người dùng, chương trình phải gọi lời gọi hệ thống lstat(2) hoặc stat(2) "
        "để nhận cấu trúc struct stat [3], [4].",
    )
    rb.add_para(
        "Trong thiết kế này, de->d_name là tên cục bộ bên trong thư mục đang mở. lstat() giải đường dẫn tương đối "
        "theo thư mục làm việc của tiến trình, không theo DIR*. Vì chương trình không chdir(), join_path() ghép "
        "thư mục cha với tên mục (ví dụ demo_tree/src_module/main.c). Một thiết kế khác có thể dùng fstatat() với descriptor thư mục cha."
    )

    rb.add_h2("2.2. Liên kết cứng (Hard link), Liên kết mềm (Symlink) và lstat() vs stat()")
    rb.add_para(
        "Hệ thống tệp POSIX hỗ trợ hai cơ chế liên kết có bản chất hoàn toàn khác nhau, được minh họa trực quan "
        "tại Hình 2.1:"
    )
    rb.add_bullet(
        "Liên kết cứng (Hard link): ",
        "Nhiều mục tên trong thư mục cùng ánh xạ trực tiếp tới một số hiệu Inode duy nhất trên cùng phân vùng. "
        "Khi tạo thêm một hard link, trường st_nlink trong Inode tăng thêm 1; dữ liệu chỉ bị xóa khỏi đĩa khi "
        "st_nlink giảm về 0 và không còn tiến trình giữ tệp mở.",
    )
    rb.add_bullet(
        "Liên kết mềm (Symbolic link – Symlink): ",
        "Là một tệp đặc biệt có số hiệu Inode riêng và kiểu tệp S_IFLNK. Nội dung dữ liệu của symlink là một "
        "chuỗi ký tự biểu diễn đường dẫn tới tệp hoặc thư mục đích. Để lấy nội dung chuỗi đích này nhằm hiển thị "
        "hậu tố \" -> target\" trong chế độ -l, chương trình phải gọi hàm readlink(2) [3].",
    )

    rb.add_figure_with_caption(
        DIAG_DIR / "diagram_2_links_inode.png",
        "Hình 2.1: Mối quan hệ giữa bảng danh mục thư mục (readdir), Inode (lstat/stat), Hard link và Symbolic link",
    )

    rb.add_para(
        "Như thể hiện ở Hình 2.1, việc lựa chọn giữa lstat() và stat() quyết định tính đúng đắn của chương trình: "
        "hàm stat() tự động đi xuyên qua symlink để đọc Inode của mục đích; nếu đích không tồn tại (broken symlink), "
        "stat() trả lỗi ENOENT khiến chương trình bỏ sót symlink đó, và nếu symlink trỏ ngược lên thư mục cha "
        "(ví dụ loop_parent -> ..), việc dùng stat() khi duyệt đệ quy -R sẽ nhận nhầm symlink là thư mục thật và "
        "gây lặp vô tận. Do đó, myls sử dụng lstat() cho toàn bộ các mục đọc từ thư mục, và chỉ sử dụng stat() "
        "để kiểm tra đích của toán hạng symlink được người dùng truyền trực tiếp trên dòng lệnh khi không bật -l "
        "hoặc khi có dấu '/' ở cuối đường dẫn [1], [2]."
    )

    rb.add_h2("2.3. Giải mã trường st_mode và số hiệu thiết bị major/minor")
    rb.add_para(
        "Trường st_mode trong struct stat có kiểu mode_t. Trên Linux, các mask truyền thống ở 16 bit thấp "
        "biểu diễn loại tệp và quyền truy cập. Mã nguồn dùng S_IS* và các mask đặt tên, không giả định độ rộng mode_t trên mọi nền tảng [1], [4]:"
    )

    mode_rows = [
        [
            "Bits 12..15 (Mặt nạ S_IFMT = 0170000)",
            "Xác định loại đối tượng tệp ở ký tự thứ 0 của chuỗi quyền hạn.",
            "S_ISREG ('-'), S_ISDIR ('d'), S_ISLNK ('l'), S_ISCHR ('c'), S_ISBLK ('b'), S_ISFIFO ('p'), S_ISSOCK ('s').",
        ],
        [
            "Bits 9..11 (S_ISUID=04000, S_ISGID=02000, S_ISVTX=01000)",
            "Ba bit đặc biệt: Set-UID, Set-GID và Sticky bit, tác động lên các vị trí ký tự thực thi thứ 3, 6 và 9.",
            "Nếu có cả bit đặc biệt và bit thực thi tương ứng: in chữ thường ('s' hoặc 't'). Nếu có bit đặc biệt nhưng không có bit thực thi: in chữ hoa ('S' hoặc 'T').",
        ],
        [
            "Bits 0..8 (S_IRWXU, S_IRWXG, S_IRWXO)",
            "9 bit quyền Đọc (r), Ghi (w), Thực thi (x) chia cho Chủ sở hữu (User), Nhóm (Group) và Người khác (Others).",
            "Phép toán bitwise AND (mode & mask) kiểm tra từng bit để điền 'r', 'w', 'x' hoặc '-'.",
        ],
    ]
    rb.add_table_with_caption(
        "Bảng 2.1: Cấu trúc các vùng bit trong trường st_mode và quy tắc chuyển đổi sang chuỗi 10 ký tự",
        ["Vùng bit trong st_mode", "Ý nghĩa kỹ thuật", "Quy tắc hiển thị trong chế độ -l"],
        mode_rows,
        [4.2, 5.2, 6.6],
    )

    rb.add_para(
        "Đối với các tệp thiết bị ký tự (Character device, S_ISCHR, ví dụ /dev/null) và thiết bị khối (Block device, "
        "S_ISBLK, ví dụ /dev/sda), trường st_size không mang ý nghĩa kích thước dữ liệu. Thay vào đó, nhân Linux "
        "lưu mã số trình điều khiển thiết bị trong trường st_rdev. Chương trình sử dụng hai macro major(st_rdev) "
        "và minor(st_rdev) trong <sys/sysmacros.h> [3] để in cặp số \"major, minor\" (ví dụ \"1, 3\" đối với /dev/null) "
        "vào đúng vị trí cột kích thước. Các macro này thuộc giao diện Linux được dùng trong đồ án; không phải API ISO C."
    )

    rb.add_h2("2.4. Tổng hợp các giao diện hệ thống và thư viện C chuẩn sử dụng")
    rb.add_para(
        "Bảng 2.2 liệt kê đầy đủ các hàm hệ thống và hàm thư viện C chuẩn được sử dụng trong mã nguồn của myls "
        "cùng cơ chế xử lý lỗi tương ứng."
    )

    api_rows = [
        ["getopt()", "<unistd.h>", "Tách và giải mã các cờ -a, -l, -R, -S; hỗ trợ dấu '--'.", "Trả về '?' khi gặp cờ lạ -> in hướng dẫn ra stderr và trả mã thoát 2."],
        ["opendir(), readdir(), closedir()", "<dirent.h>", "Mở thư mục, đọc từng mục struct dirent và đóng directory stream.", "Đặt errno = 0 trước readdir() để tách biệt EOF với lỗi đọc; kiểm tra giá trị trả về của closedir()."],
        ["lstat(), stat()", "<sys/stat.h>", "Thu thập siêu dữ liệu Inode vào struct stat.", "Nếu trả về -1: in thông báo strerror(errno) ra stderr và nâng mã thoát (1 hoặc 2)."],
        ["readlink()", "<unistd.h>", "Đọc đường dẫn đích của symbolic link vào bộ đệm cấp phát động.", "Tự gán ký tự kết thúc chuỗi '\\0' tại vị trí len; nếu lỗi vẫn giữ metadata và nâng mã thoát."],
        ["getpwuid(), getgrgid()", "<pwd.h>, <grp.h>", "Tra cứu tên người dùng và tên nhóm từ st_uid, st_gid.", "Nếu trả về NULL (UID/GID không có tên): tự động in mã số nguyên để tránh lỗi con trỏ NULL."],
        ["localtime_r()", "<time.h>", "Chuyển đổi dấu thời gian st_mtime sang cấu trúc lịch cục bộ an toàn luồng.", "Kiểm tra khoảng cách 6 tháng (15.778.476 giây) so với time(NULL) để chọn định dạng Giờ:Phút hoặc Năm."],
        ["qsort(), strcmp()", "<stdlib.h>, <string.h>", "Sắp xếp mảng động FileEntry theo kích thước (-S) hoặc tên từ điển.", "Dùng phép so sánh > và < trên kiểu off_t (64 bit trên môi trường x86_64 đã thử), kết hợp strcmp() khi kích thước bằng nhau."],
        ["fflush(), ferror()", "<stdio.h>", "Đẩy bộ đệm đầu ra và kiểm tra lỗi ghi thiết bị trên stdout.", "Phát hiện lỗi ghi như ENOSPC trên /dev/full, báo lỗi ra stderr và trả mã thoát 2."],
    ]
    rb.add_table_with_caption(
        "Bảng 2.2: Danh mục các API chuẩn POSIX và C11 sử dụng trong chương trình myls",
        ["Hàm API", "Header", "Vai trò trong chương trình", "Cơ chế kiểm soát lỗi"],
        api_rows,
        [3.3, 2.5, 4.8, 5.4],
    )

    # --------------------------------------------------------------------------
    # CHƯƠNG 3: THIẾT KẾ KIẾN TRÚC VÀ CÀI ĐẶT CHƯƠNG TRÌNH
    # --------------------------------------------------------------------------
    rb.add_chapter_heading("CHƯƠNG 3. THIẾT KẾ KIẾN TRÚC VÀ CÀI ĐẶT CHƯƠNG TRÌNH", new_page=True)

    rb.add_h2("3.1. Kiến trúc phân tách 5 module và luồng dữ liệu tổng thể")
    rb.add_para(
        "Chương trình được thiết kế theo kiến trúc 5 module độc lập, giao tiếp qua tệp tiêu đề trung tâm "
        "include/myls.h. Luồng đi của dữ liệu từ khi nhận tham số dòng lệnh argv[] cho đến khi kết xuất và "
        "duyệt đệ quy được trình bày tại Hình 3.1 và phân công trách nhiệm tại Bảng 3.1."
    )

    rb.add_figure_with_caption(
        DIAG_DIR / "diagram_1_pipeline.png",
        "Hình 3.1: Sơ đồ luồng thực thi tổng thể từ phân tích tham số dòng lệnh đến kết xuất và đệ quy",
    )

    mod_rows = [
        ["include/myls.h", "Định nghĩa 4 cấu trúc dữ liệu (LsOptions, FileEntry, EntryList, ColumnWidths), các hằng mã thoát (0, 1, 2), hàm nội tuyến elevate_status() và nguyên mẫu hàm."],
        ["src/options.c", "Cài đặt parse_options() và print_usage(): giải mã tổ hợp cờ -a, -l, -R, -S bằng getopt(), xử lý cờ --help và dấu phân cách '--'."],
        ["src/main.c", "Điều phối chính: phân loại toán hạng CLI thành file_list và dir_list, xử lý ngoại lệ symlink tới thư mục trên CLI, gọi hiển thị 2 pha và kiểm tra lỗi flush stdout."],
        ["src/directory.c", "Quản lý mảng động EntryList có kiểm tra tràn SIZE_MAX, ghép đường dẫn join_path(), đọc metadata qua lstat()/readlink() và duyệt đệ quy process_directory()."],
        ["src/sorting.c", "Cài đặt sort_entries() cùng 2 hàm so sánh cmp_by_name_asc() và cmp_by_size_desc() cho qsort()."],
        ["src/display.c", "Giải mã chuỗi quyền 10 ký tự format_mode_string(), tính độ rộng cột 2 lượt compute_column_widths(), làm tròn dòng total 1 KiB và kiểm tra lỗi ghi stdout."],
    ]
    rb.add_table_with_caption(
        "Bảng 3.1: Phân công trách nhiệm giữa các tệp mã nguồn trong dự án",
        ["Tệp nguồn", "Trách nhiệm kỹ thuật chính"],
        mod_rows,
        [3.8, 12.2],
    )

    rb.add_h2("3.2. Cấu trúc dữ liệu cốt lõi và quản lý quyền sở hữu bộ nhớ (Ownership)")
    rb.add_para(
        "Để không bị giới hạn cứng số lượng tệp trong một thư mục và tiết kiệm bộ nhớ khi duyệt cây đệ quy sâu, "
        "cấu trúc FileEntry và EntryList trong include/myls.h được thiết kế với quyền sở hữu bộ nhớ rõ ràng:"
    )

    rb.add_code_snippet(
        "Trích đoạn định nghĩa cấu trúc FileEntry và EntryList (include/myls.h):",
        """typedef struct {
    char *name;        /* Tên hiển thị (cấp phát động trên Heap) */
    char *full_path;   /* Đường dẫn đầy đủ để gọi lstat/readlink/đệ quy */
    struct stat st;    /* Siêu dữ liệu thu được từ lstat()/stat() */
    bool stat_ok;      /* Đánh dấu đọc metadata thành công */
    char *link_target; /* Chỉ cấp phát động khi là Symlink ở chế độ -l; ngược lại là NULL */
} FileEntry;

typedef struct {
    FileEntry *items;  /* Mảng động các FileEntry */
    size_t count;      /* Số phần tử hiện có */
    size_t capacity;   /* Sức chứa hiện tại (khởi đầu 64, nhân đôi khi đầy) */
} EntryList;""",
    )

    rb.add_para(
        "Thay vì đặt một mảng tĩnh char link_target[PATH_MAX] (chiếm cố định 4096 bytes) bên trong mọi phần tử "
        "FileEntry, chương trình sử dụng con trỏ char *link_target và đặt bằng NULL cho toàn bộ các tệp thông thường. "
        "Vùng nhớ cho link_target chỉ được cấp phát khi mục đó thực sự là symbolic link (S_ISLNK) và người dùng "
        "bật cờ -l. Đồng thời, trong hàm entry_list_push() (src/directory.c), trước khi nhân đôi capacity hoặc "
        "tính kích thước truyền vào realloc(), chương trình kiểm tra điều kiện tràn số nguyên:"
    )

    rb.add_code_snippet(
        "Trích đoạn kiểm tra tràn số nguyên và hủy mục cấp phát dở dang (src/directory.c):",
        """if (list->count >= list->capacity) {
    size_t new_cap = (list->capacity == 0) ? INITIAL_LIST_CAPACITY : (list->capacity * 2);
    if (list->capacity > SIZE_MAX / 2 || new_cap > SIZE_MAX / sizeof(FileEntry)) {
        errno = ENOMEM;
        fprintf(stderr, "myls: memory allocation failed\\n");
        return false;
    }
    FileEntry *new_items = (FileEntry *)realloc(list->items, new_cap * sizeof(FileEntry));
    if (new_items == NULL) {
        fprintf(stderr, "myls: memory allocation failed\\n");
        return false;
    }
    list->items = new_items;
    list->capacity = new_cap;
}""",
    )

    rb.add_h2("3.3. Thuật toán duyệt đệ quy (-R), chốt chặn vòng lặp và đánh đổi tài nguyên")
    rb.add_para(
        "Khi bật tùy chọn -R, hàm process_directory() trong src/directory.c thực hiện quy trình 4 bước tuần tự: "
        "(1) Mở thư mục bằng opendir() và đọc toàn bộ các mục thỏa điều kiện lọc vào mảng động EntryList list; "
        "(2) Đóng ngay luồng thư mục bằng closedir(dir) và kiểm tra lỗi; "
        "(3) Sắp xếp list bằng sort_entries() và kết xuất toàn bộ mục của thư mục hiện tại ra màn hình; "
        "(4) Duyệt qua list theo đúng thứ tự đã sắp xếp và gọi đệ quy process_directory() cho các mục con thỏa "
        "hàm kiểm tra is_recursible_subdir(). Hình 3.2 minh họa thứ tự duyệt cây và quy tắc bỏ qua . / .. / symlink khi đi xuống."
    )

    rb.add_figure_with_caption(
        DIAG_DIR / "diagram_3_recursion_tree.png",
        "Hình 3.2: Thứ tự duyệt -aR và quy tắc bỏ qua . / .. / symlink khi đi xuống cây",
    )

    rb.add_code_snippet(
        "Trích đoạn kiểm tra điều kiện đệ quy và đọc thư mục an toàn (src/directory.c):",
        """static bool is_recursible_subdir(const FileEntry *entry)
{
    if (!entry->stat_ok || !S_ISDIR(entry->st.st_mode)) {
        return false; /* Chặn mọi tệp thường và Symbolic Link (nhờ dùng lstat) */
    }
    if (strcmp(entry->name, ".") == 0 || strcmp(entry->name, "..") == 0) {
        return false; /* Chặn '.' và '..' khi chạy kết hợp -aR */
    }
    return true;
}""",
    )

    rb.add_para(
        "Phân tích sự đánh đổi giữa File Descriptor (FD) và Bộ nhớ (RAM): Việc gọi closedir(dir) ngay trước vòng "
        "lặp đệ quy giúp giải phóng File Descriptor của thư mục hiện tại trước khi mở thư mục con. Nhờ vậy, số "
        "lượng FD dùng cho việc đọc thư mục tại mọi thời điểm chỉ là 1 FD, không bị chạm ngưỡng giới hạn RLIMIT_NOFILE "
        "(thường là 1024) khi cây thư mục sâu. Tuy nhiên, cần lưu ý sự đánh đổi về bộ nhớ: vì lời gọi đệ quy nằm "
        "trước lệnh entry_list_free(&list), các danh sách EntryList của những thư mục cha dọc theo nhánh đệ quy "
        "hiện tại vẫn được giữ trong bộ nhớ cho tới khi các thư mục con của nhánh đó duyệt xong. Do đó, bộ nhớ "
        "tức thời tỷ lệ với tổng số mục trên các mức từ gốc đến nút đang xét, chứ không chỉ riêng thư mục hiện hành."
    )

    rb.add_h2("3.4. Thuật toán sắp xếp (-S) và định dạng kết xuất hai lượt (-l)")
    rb.add_para(
        "Trong src/sorting.c, hàm cmp_by_size_desc() so sánh trường st_size (kiểu off_t (64 bit trên môi trường x86_64 đã thử) có dấu) bằng hai "
        "phép so sánh quan hệ (a->st.st_size > b->st.st_size trả về -1; < trả về 1) thay vì lấy hiệu trực tiếp, "
        "nhằm loại bỏ hoàn toàn nguy cơ tràn số nguyên 32-bit khi so sánh các tệp lớn hơn 2 GiB. Khi hai mục có "
        "cùng kích thước, hàm gọi strcmp(a->name, b->name) để phân định thứ tự tăng dần theo tên."
    )
    rb.add_para(
        "Trong src/display.c, chế độ -l sử dụng thuật toán hai lượt (Two-pass alignment) để căn chỉnh các cột "
        "thẳng hàng và định dạng đầy đủ 7 trường dữ liệu trên mỗi dòng như minh họa ở Hình 3.3:"
    )

    rb.add_figure_with_caption(
        DIAG_DIR / "diagram_4_long_format.png",
        "Hình 3.3: Giải phẫu các trường thông tin trên một dòng kết xuất ở chế độ -l",
    )

    rb.add_bullet(
        "Lượt 1 (Tính độ rộng cột và tổng khối 1 KiB): ",
        "Hàm compute_column_widths() duyệt qua danh sách để tìm số chữ số lớn nhất của st_nlink, độ dài chuỗi "
        "tên owner/group và độ rộng cột kích thước. Nếu danh sách có chứa thiết bị ký tự hoặc khối (S_ISCHR / S_ISBLK), "
        "hàm tính độ rộng lớn nhất của major(st_rdev) và minor(st_rdev), sau đó đồng bộ độ rộng tổng \"major, minor\" "
        "với độ rộng st_size của tệp thường. Tổng số khối 512-byte (total_blocks_512) được đổi sang đơn vị 1 KiB "
        "có làm tròn lên theo công thức: total_1k = (total_blocks_512 + 1) / 2.",
    )
    rb.add_bullet(
        "Lượt 2 (Kết xuất và kiểm tra lỗi ghi stdout): ",
        "Chương trình in từng dòng theo độ rộng động đã tính. Mọi thao tác printf() và lệnh fflush(stdout) ở cuối "
        "chương trình đều được kiểm tra lỗi qua flush_and_check_stdout(). Nếu luồng đầu ra gặp sự cố (chẳng hạn "
        "khi chuyển hướng ra thiết bị đầy /dev/full gây lỗi ENOSPC), chương trình ghi thông báo lỗi ra stderr và "
        "trả về mã thoát nghiêm trọng 2.",
    )

    # --------------------------------------------------------------------------
    # CHƯƠNG 4: KIỂM THỬ THỰC NGHIỆM VÀ ĐÁNH GIÁ KẾT QUẢ
    # --------------------------------------------------------------------------
    rb.add_chapter_heading("CHƯƠNG 4. KIỂM THỬ THỰC NGHIỆM VÀ ĐÁNH GIÁ KẾT QUẢ", new_page=True)

    rb.add_h2("4.1. Môi trường thực nghiệm và thiết kế cây dữ liệu kiểm thử (Fixtures)")
    rb.add_para(
        f"Toàn bộ quá trình biên dịch và kiểm thử được thực hiện trên môi trường {env_info.get('OS_NAME', 'Ubuntu 24.04 LTS')} "
        f"(nhân {env_info.get('KERNEL_VER', 'Linux WSL2')}), trình biên dịch {env_info.get('GCC_VER', 'GCC 13.3.0')} "
        f"và đối chiếu trực tiếp với {env_info.get('COREUTILS_VER', 'GNU coreutils 9.4')}. "
        "Một quyết định kỹ thuật quan trọng trong kịch bản kiểm thử (tests/run_tests.sh) là khởi tạo toàn bộ cây "
        "dữ liệu mẫu trên phân vùng /tmp của Linux thay vì thư mục /mnt/p (nơi gắn ổ đĩa Windows qua trình điều "
        "khiển DrvFs). Quyền trên ổ Windows phụ thuộc quyền Windows, tùy chọn mount và chế độ metadata; không nên "
        "giả định chmod hoặc các tệp đặc biệt có hành vi như filesystem Linux. Fixture /tmp giúp đánh giá nhất quán [5]."
    )

    rb.add_h2("4.2. Phương pháp kiểm thử và độ bao phủ kịch bản")
    rb.add_para(
        f"Bộ kiểm thử tự động gồm tổng cộng {reg_suite['total']} kịch bản, được chia thành 3 nhóm kiểm chứng "
        "chặt chẽ cả đầu ra chuẩn (stdout diff -u), luồng báo lỗi (stderr) và mã thoát (exit code):"
    )
    rb.add_bullet(
        "Nhóm 1 – Toàn bộ 16 tổ hợp của 4 cờ (Kịch bản 01 đến 16): ",
        "Kiểm tra vét cạn 2^4 = 16 tổ hợp của {-a, -l, -R, -S} (từ không cờ đến -alRS) trên cây thư mục đa tầng "
        "có chứa tệp ẩn, thư mục ẩn, tệp trùng kích thước, sparse file lớn hơn 2 GiB (2,5 GB), hard link, symlink "
        "tới tệp/thư mục, broken symlink và symlink trỏ ngược lên thư mục cha (..).",
    )
    rb.add_bullet(
        "Nhóm 2 – Cú pháp cờ, tệp đặc biệt và unit test nội bộ (Kịch bản 17 đến 27): ",
        "Kiểm tra viết rời và đảo thứ tự cờ (-S -R -l -a), thư mục rỗng, thư mục có hơn 64 mục (80 tệp để đi qua "
        "nhánh nhân đôi capacity 64 -> 128), tên tệp bắt đầu bằng '-' đi kèm '--', FIFO, Unix domain socket, "
        "toán hạng symlink trên CLI (không có -l, có -l, và có dấu '/' cuối), thiết bị ký tự /dev/null trộn với "
        "tệp thường, cùng bộ unit test C nội bộ (tests/test_internals.c) kiểm tra định dạng block device tổng hợp, "
        "làm tròn khối 1 KiB và bảo vệ tràn số SIZE_MAX.",
    )
    rb.add_bullet(
        "Nhóm 3 – Đường lỗi, phân quyền, /dev/full và tiêm lỗi Fault Injection (Kịch bản 28 đến 37): ",
        "Kiểm tra mã thoát 2 cho đối số CLI không tồn tại, broken symlink có dấu '/' cuối, thư mục CLI bị khóa "
        "quyền đọc (chmod 0000), cờ không hợp lệ (-z) và lỗi ghi stdout (/dev/full); kiểm tra mã thoát 1 khi gặp "
        "thư mục con bị khóa quyền đọc trong lúc đệ quy -R; kiểm tra giữ mã lỗi cao nhất (2) khi trộn lỗi CLI với "
        "lỗi thư mục con; và sử dụng thư viện LD_PRELOAD (tests/fault_inject.c) để tiêm lỗi EIO/ENOMEM vào các hàm "
        "readdir(), readlink(), closedir() và realloc().",
    )

    # Bảng tổng hợp kết quả thực nghiệm lấy trực tiếp từ test_results.json
    summary_rows = [
        [
            "Kiểm thử chức năng & đường lỗi (make test)",
            f"./myls ({reg_suite['command']})",
            f"{reg_suite['passed']} / {reg_suite['total']} đạt ({reg_suite['failed']} lỗi, {reg_suite['skipped']} bỏ qua)",
            f"Exit code = {reg_suite['exit_code']}. 26 ca đối chiếu GNU ls (stdout và mã thoát); 11 ca kiểm hợp đồng, helper nội bộ hoặc lỗi giả lập.",
        ],
        [
            "Kiểm thử bộ nhớ & hành vi bất định (make sanitize)",
            f"./myls_sanitize ({san_suite['command']})",
            f"{san_suite['passed']} / {san_suite['total']} đạt ({san_suite['failed']} lỗi, {san_suite['skipped']} bỏ qua)",
            f"Exit code = {san_suite['exit_code']}. {san_suite['statement']}",
        ],
        [
            "Kịch bản trình diễn bảo vệ (make demo)",
            "./myls (tests/run_demo.sh)",
            "Hoàn tất 4 bước demo trên fixture cô lập /tmp",
            "Các bước mặc định, -a, -lS, -alRS trả mã 0; bước thử đối số lỗi trả mã 2 đúng hợp đồng.",
        ],
    ]
    rb.add_table_with_caption(
        "Bảng 4.1: Tổng hợp kết quả thực nghiệm từ tệp dữ liệu evidence/test_results.json",
        ["Hạng mục thực nghiệm", "Lệnh / Binary thực thi", "Kết quả ghi nhận", "Đánh giá chi tiết"],
        summary_rows,
        [3.8, 3.6, 3.4, 5.2],
    )

    rb.add_h2("4.3. Minh chứng đầu ra thực nghiệm từ log kiểm thử và demo")
    rb.add_para(
        "Các Hình 4.1, Hình 4.2 và Hình 4.3 dưới đây trích các dòng đầu ra thực nghiệm đã được ghi nhận "
        "trong các tệp log tại thư mục evidence/ (evidence/test_suite_report.txt và evidence/demo_session_output.txt)."
    )

    rb.add_figure_with_caption(
        EVIDENCE_DIR / "test_suite_capture.png",
        f"Hình 4.1: Đầu ra thực nghiệm trình bày lại từ log kiểm thử tự động ({reg_suite['passed']}/{reg_suite['total']} kịch bản đạt)",
        width_cm=15.6,
    )

    rb.add_figure_with_caption(
        EVIDENCE_DIR / "demo_basic_flags_capture.png",
        "Hình 4.2: Đầu ra thực nghiệm trình bày lại từ log demo cho các tùy chọn mặc định, -a và -lS",
        width_cm=15.6,
    )

    rb.add_figure_with_caption(
        EVIDENCE_DIR / "demo_combined_alRS_capture.png",
        "Hình 4.3: Đầu ra thực nghiệm trình bày lại từ log demo cho tổ hợp -alRS và xử lý đối số lỗi",
        width_cm=15.6,
    )

    # --------------------------------------------------------------------------
    # CHƯƠNG 5: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN
    # --------------------------------------------------------------------------
    rb.add_chapter_heading("CHƯƠNG 5. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", new_page=True)

    rb.add_h2("5.1. Kết quả đạt được")
    rb.add_para(
        "Đồ án đã hoàn thành việc nghiên cứu cơ chế quản lý thư mục và siêu dữ liệu của hệ thống tệp Linux, đồng "
        "thời xây dựng thành công chương trình myls bằng ngôn ngữ C (chuẩn C11 + POSIX.1-2008) đáp ứng các yêu cầu "
        "kỹ thuật của Đề tài 17:"
    )
    rb.add_bullet(
        "Hoàn thiện đầy đủ bốn tùy chọn -a, -l, -R, -S và mọi tổ hợp kết hợp: ",
        "Chương trình giải mã chính xác 10 ký tự loại tệp và quyền truy cập (bao gồm SUID, SGID, Sticky bit), "
        "hiển thị số hiệu thiết bị major, minor cho character/block device, căn lề cột động hai lượt, sắp xếp "
        "an toàn trên kiểu off_t và bỏ qua . / .. / symlink khi đi xuống cây. Chưa có phát hiện chu trình do bind mount.",
    )
    rb.add_bullet(
        "Quản lý tài nguyên và xử lý lỗi chặt chẽ: ",
        "Kiểm soát đầy đủ các đường lỗi của opendir(), readdir() (đặt errno = 0), lstat(), readlink(), closedir(), "
        "realloc() và lỗi ghi luồng stdout (fflush/ferror); phân biệt chính xác mã thoát 0, 1 và 2.",
    )
    rb.add_bullet(
        "Kiểm chứng thực nghiệm có khả năng tái lập: ",
        f"Bản thường đạt {reg_suite['passed']}/{reg_suite['total']} ca, bỏ qua {reg_suite['skipped']}; bản sanitizer đạt {san_suite['passed']}/{san_suite['total']} ca, bỏ qua {san_suite['skipped']}. Không phát hiện lỗi bởi ASan/LSan/UBSan trong các ca đã chạy; kết quả không chứng minh mọi đầu vào đều không có lỗi.",
    )

    rb.add_h2("5.2. Hạn chế hiện tại và hướng phát triển")
    rb.add_para(
        "Bên cạnh các kết quả đã đạt được trong phạm vi đề tài, chương trình vẫn còn một số giới hạn kỹ thuật "
        "và có thể được mở rộng theo thứ tự ưu tiên sau:"
    )
    rb.add_bullet(
        "Ưu tiên 1 – Phát hiện chu trình hệ thống tệp theo cặp (st_dev, st_ino): ",
        "Bổ sung bảng băm hoặc cây tìm kiếm lưu cặp (st_dev, st_ino) của các thư mục tổ tiên đang mở trên nhánh "
        "đệ quy để phát hiện chu trình do bind mount hoặc hard link thư mục đặc biệt.",
    )
    rb.add_bullet(
        "Ưu tiên 2 – Sử dụng họ hàm *at (openat, fstatat, readlinkat): ",
        "Thay thế việc nối chuỗi đường dẫn đầy đủ bằng việc truyền descriptor của thư mục cha vào fstatat() và "
        "readlinkat(), giúp tăng hiệu năng tra cứu đường dẫn sâu và giảm rủi ro khi cây thư mục bị đổi tên đồng thời.",
    )
    rb.add_bullet(
        "Ưu tiên 3 – Hỗ trợ dàn nhiều cột trên terminal và làm sạch ký tự điều khiển: ",
        "Khi đầu ra là thiết bị đầu cuối (isatty(STDOUT_FILENO)), truy vấn kích thước cửa sổ bằng ioctl(TIOCGWINSZ) "
        "để dàn tên tệp thành nhiều cột và thay thế các ký tự không in được trong tên tệp bằng dấu '?' để bảo vệ terminal.",
    )

    # --------------------------------------------------------------------------
    # TÀI LIỆU THAM KHẢO
    # --------------------------------------------------------------------------
    sec_back = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_a4_margins(sec_back)
    set_section_page_number_format(sec_back, fmt="lowerRoman", start=1)
    configure_header_footer(sec_back, "Tài liệu tham khảo và phụ lục", True)
    rb.add_chapter_heading("TÀI LIỆU THAM KHẢO", new_page=False)

    refs = [
        "[1] IEEE and The Open Group, \"The Open Group Base Specifications Issue 7, 2018 edition (POSIX.1-2008 / IEEE Std 1003.1-2017) – ls utility, sys/stat.h, dirent.h,\" 2018. [Online]. Available: https://pubs.opengroup.org/onlinepubs/9699919799/",
        "[2] Free Software Foundation, \"GNU Coreutils Manual – Section 10.1: ls: List directory contents,\" GNU Project, 2024. [Online]. Available: https://www.gnu.org/software/coreutils/manual/html_node/ls-invocation.html",
        "[3] Michael Kerrisk et al., \"Linux Programmer's Manual (man-pages): stat(2), lstat(2), readlink(2), opendir(3), readdir(3), closedir(3), getopt(3), sysmacros(7),\" 2024. [Online]. Available: https://man7.org/linux/man-pages/",
        "[4] Michael Kerrisk, The Linux Programming Interface: A Linux and UNIX System Programming Handbook, San Francisco, CA: No Starch Press, 2010, Chapters 14, 15, and 18.",
        "[5] Microsoft Corporation, \"File Permissions for WSL (Windows Subsystem for Linux DrvFs and Metadata),\" Microsoft Learn Documentation, 2024. [Online]. Available: https://learn.microsoft.com/en-us/windows/wsl/file-permissions",
        "[6] Trường Đại học Công nghệ Thông tin và Truyền thông Việt – Hàn (VKU), \"Quyết định số 675/QĐ-ĐHVH ban hành Quy định triển khai đồ án, khóa luận tốt nghiệp (Phụ lục Mẫu 3 – Quy cách trình bày báo cáo),\" Đà Nẵng, ngày 03/04/2024. Tham khảo quy cách hình thức cho báo cáo học phần; không áp dụng yêu cầu độ dài của khóa luận. [Online]. Available: https://daotao.vku.udn.vn/uploads/daotao/675-qd-ban-hanh-quy-dinh-trien-khai-do-an-khoa-luan-tot-nghiep-trong-cac-ctdt-dh-cua-vku-0001.pdf",
    ]
    for ref in refs:
        p_ref = doc.add_paragraph(style="Normal")
        p_ref.paragraph_format.left_indent = Cm(0.8)
        p_ref.paragraph_format.first_line_indent = Cm(-0.8)
        p_ref.paragraph_format.space_after = Pt(6)
        r = p_ref.add_run(ref)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12.5)
        r.font.color.rgb = BLACK

    # --------------------------------------------------------------------------
    # PHỤ LỤC
    # --------------------------------------------------------------------------
    rb.add_chapter_heading("PHỤ LỤC. DANH SÁCH KỊCH BẢN KIỂM THỬ TỰ ĐỘNG VÀ LỆNH TÁI LẬP", new_page=True)

    rb.add_h2("Phụ lục A. Quy trình lệnh kiểm chứng độc lập trên Ubuntu WSL2")
    rb.add_code_snippet(
        "Các lệnh biên dịch, chạy kiểm thử thường, kiểm thử Sanitizer và chạy demo:",
        """# 1. Biên dịch chương trình với bộ cờ cảnh báo nghiêm ngặt
make clean && make

# 2. Chạy bộ kiểm thử tự động đối chiếu với GNU ls trên /tmp
make test

# 3. Biên dịch và chạy toàn bộ bộ kiểm thử với ASan + UBSan + LSan
make sanitize

# 4. Chạy kịch bản demo 4 bước trên cây dữ liệu mẫu cô lập tại /tmp
make demo""",
    )

    rb.add_h2(f"Phụ lục B. Danh sách chi tiết {reg_suite['total']} kịch bản trong bộ kiểm thử tự động")
    case_rows = [
        [f"Test {c['id']:02d}", c["description"], c["status"]]
        for c in reg_suite["cases"]
    ]
    rb.add_table_with_caption(
        f"Bảng P.1: Danh sách toàn bộ {reg_suite['total']} kịch bản kiểm thử tự động và trạng thái thực thi",
        ["Mã kịch bản", "Nội dung kiểm chứng", "Trạng thái"],
        case_rows,
        [2.5, 11.3, 2.2],
    )

    # Lưu ra file tạm trước, sau đó dùng Word COM để cập nhật TOC/fields và xuất PDF
    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    tmp_docx = BUILD_DIR / "BaoCao_tmp.docx"
    tmp_pdf = BUILD_DIR / "BaoCao_tmp.pdf"
    if tmp_docx.exists():
        tmp_docx.unlink()
    if tmp_pdf.exists():
        tmp_pdf.unlink()

    doc.save(str(tmp_docx))
    export_and_update_via_word(tmp_docx, tmp_pdf)

    if not tmp_pdf.exists() or tmp_pdf.stat().st_size == 0:
        raise RuntimeError("PDF export failed: temporary PDF was not created.")

    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(tmp_docx, out_docx)
    shutil.copy2(tmp_pdf, out_pdf)
    print(f"[OK] Generated VKU report: {out_docx} and {out_pdf}; visual inspection required.")


def export_and_update_via_word(docx_path: Path, pdf_path: Path) -> None:
    subprocess.run(["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
        str(PROJECT_ROOT / "scripts" / "export_office.ps1"), "-Kind", "Word",
        "-InputPath", str(docx_path.resolve()), "-OutputPath", str(pdf_path.resolve())], check=True)


if __name__ == "__main__":
    build_report(DOCS_DIR / "BaoCao.docx", DOCS_DIR / "BaoCao.pdf")
